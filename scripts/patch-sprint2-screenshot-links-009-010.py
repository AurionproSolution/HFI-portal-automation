#!/usr/bin/env python3
"""Add missing screenshot hyperlinks for US-DLR-009 and US-DLR-010 PASS rows only."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "reports" / "Sprint2-Management-Regression-Report-Aug15-16.xlsx"
MANIFEST = ROOT / "reports" / "sprint2-pass-manifest.json"
EVIDENCE_DIR = ROOT / "reports" / "sprint2-evidence"
STORIES = ("US-DLR-009", "US-DLR-010")
ALL_SCREENSHOT_COL = 25
STORY_SCREENSHOT_COL = 9


def set_file_hyperlink(cell, rel_path: str) -> bool:
    rel_path = (rel_path or "").strip()
    if not rel_path:
        return False
    first = rel_path.split(";")[0].strip()
    target = (ROOT / first).resolve()
    if not target.exists():
        return False
    cell.value = first
    try:
        link_target = target.relative_to(XLSX.parent.resolve()).as_posix()
    except ValueError:
        link_target = target.as_uri()
    cell.hyperlink = link_target
    cell.font = Font(color="0563C1", underline="single")
    return True


def evidence_path(story_id: str, tc_id: str) -> str:
    return f"reports/sprint2-evidence/{story_id}-{tc_id}-PASS.png"


def main() -> None:
    if not XLSX.exists():
        raise SystemExit(f"Missing {XLSX}")
    if not MANIFEST.exists():
        raise SystemExit(f"Missing {MANIFEST}")

    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    pass_by_story = {story: set(manifest.get(story, [])) for story in STORIES}

    wb = load_workbook(XLSX)
    updated_all = 0
    updated_story = 0
    skipped_existing = 0
    missing_files = []

    ws_all = wb["All Test Cases"]
    for row in ws_all.iter_rows(min_row=2, values_only=False):
        story = row[1].value
        if story not in STORIES:
            continue
        tc = row[4].value
        status = row[15].value
        if status != "PASS" or tc not in pass_by_story[story]:
            continue

        cell = row[ALL_SCREENSHOT_COL - 1]
        if (cell.value or "").strip():
            skipped_existing += 1
            continue

        shot = evidence_path(story, tc)
        if set_file_hyperlink(cell, shot):
            updated_all += 1
        else:
            missing_files.append(shot)

    for story in STORIES:
        sheet_name = story.replace("-", "_")
        if sheet_name not in wb.sheetnames:
            raise SystemExit(f"Missing sheet {sheet_name}")
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=False):
            tc = row[0].value
            status = row[2].value
            if status != "PASS" or tc not in pass_by_story[story]:
                continue

            cell = row[STORY_SCREENSHOT_COL - 1]
            if (cell.value or "").strip():
                skipped_existing += 1
                continue

            shot = evidence_path(story, tc)
            if set_file_hyperlink(cell, shot):
                updated_story += 1
            else:
                missing_files.append(shot)

    wb.save(XLSX)
    wb.close()

    print(f"Updated All Test Cases links: {updated_all}")
    print(f"Updated per-story sheet links: {updated_story}")
    print(f"Skipped existing links: {skipped_existing}")
    if missing_files:
        unique_missing = sorted(set(missing_files))
        print(f"Missing evidence files ({len(unique_missing)}):")
        for path in unique_missing:
            print(f"  {path}")


if __name__ == "__main__":
    main()
