#!/usr/bin/env python3
"""Verify Sprint 2 Excel screenshot hyperlinks point to existing files."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "reports" / "Sprint2-Management-Regression-Report-Aug15-16.xlsx"
SCREENSHOT_COL = 25


def main() -> None:
    if not XLSX.exists():
        raise SystemExit(f"Missing {XLSX}")

    wb = load_workbook(XLSX)
    ws = wb["All Test Cases"]
    ok = 0
    empty_pass = 0
    broken = 0
    rows = []

    for row in ws.iter_rows(min_row=2, values_only=False):
        story = row[1].value
        tc = row[4].value
        status = row[15].value
        cell = row[SCREENSHOT_COL - 1]
        path = (cell.value or "").strip()
        link = cell.hyperlink.target if cell.hyperlink else ""

        if status != "PASS":
            continue
        if story == "UC-ADM-001":
            continue
        if not path:
            empty_pass += 1
            rows.append((story, tc, "EMPTY", ""))
            continue

        if link and not link.startswith("file:"):
            target = (XLSX.parent / link).resolve()
        elif link and link.startswith("file:"):
            target = Path(link.replace("file:///", "").replace("file://", ""))
        else:
            target = (ROOT / path).resolve()

        exists = target.exists()
        if exists:
            ok += 1
        else:
            broken += 1
        rows.append((story, tc, "OK" if exists else "BROKEN", str(target)))

    wb.close()
    print(f"PASS rows (6 stories): links ok={ok} empty={empty_pass} broken={broken}")
    for story, tc, state, target in rows:
        if state != "OK":
            print(f"  {state} {story} {tc} -> {target}")


if __name__ == "__main__":
    main()
