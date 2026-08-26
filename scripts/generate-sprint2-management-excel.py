#!/usr/bin/env python3
"""
Sprint 2 management Excel — consolidates TC catalogs, Sat/Sun execution logs,
markdown reports, and test-result evidence paths.
"""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
TEST_RESULTS = ROOT / "test-results"
SPRINT2_EVIDENCE_DIR = REPORTS / "sprint2-evidence"
PASS_MANIFEST = REPORTS / "sprint2-pass-manifest.json"
OUT_XLSX = REPORTS / "Sprint2-Management-Regression-Report-Aug15-16.xlsx"

SAT = "2026-08-15"
SUN = "2026-08-16"

SUITES = [
    {
        "story_id": "UC-ADM-001",
        "story_name": "Admin Console (ADM)",
        "portal": "Admin / OEM Console",
        "total": 55,
        "catalog": ("json_adm", REPORTS / "uc-adm-001-admin-console-cases.json"),
        "report_md": REPORTS / "admin-console-regression-execution-report.md",
        "logs": [
            REPORTS / "admin-console-tc001-005-run.log",
            REPORTS / "admin-console-tc001-005-run2.log",
            REPORTS / "admin-console-tc001-005-run3.log",
            REPORTS / "admin-console-tc006-010-run1.log",
            REPORTS / "admin-console-tc006-010-run2.log",
            REPORTS / "admin-console-tc006-010-run3.log",
            REPORTS / "admin-console-tc006-010-run4.log",
            REPORTS / "admin-console-tc011-020-run1.log",
            REPORTS / "admin-console-tc019-020-run2.log",
            REPORTS / "admin-console-tc021-030-run1.log",
            REPORTS / "admin-console-tc021-030-run2.log",
            REPORTS / "admin-console-tc021-030-run3.log",
            REPORTS / "admin-console-tc025-030-run3.log",
            REPORTS / "admin-console-tc031-040-run1.log",
            REPORTS / "admin-console-tc031-040-run2.log",
            REPORTS / "admin-console-tc035-040-run3.log",
            REPORTS / "admin-console-tc041-055-run1.log",
            REPORTS / "admin-console-tc041-055-run2.log",
            REPORTS / "admin-console-login-comms-banner-sprint-run.log",
            REPORTS / "admin-console-regression-dev-run1.log",
        ],
        "dealer": "ved.prakash (OEM Admin)",
        "run_dates": f"{SAT} (slots/batches)",
    },
    {
        "story_id": "US-OEM-002",
        "story_name": "Dealer Limits (OEM2)",
        "portal": "OEM Portal",
        "total": 35,
        "catalog": ("json_oem2", REPORTS / "us-oem-002-dealer-limits-cases.json"),
        "report_md": REPORTS / "dealer-limits-regression-execution-report.md",
        "logs": [
            REPORTS / "dealer-limits-slot1-run.log",
            REPORTS / "dealer-limits-slot2-run.log",
            REPORTS / "dealer-limits-checker-run.log",
            REPORTS / "dealer-limits-checker-run2.log",
            REPORTS / "dealer-limits-regression-dev-run.log",
            REPORTS / "dealer-limits-regression-dev-run2.log",
            REPORTS / "dealer-limits-regression-rerun-automation-fix.log",
        ],
        "dealer": "ved.prakash (OEM Checker)",
        "run_dates": SAT,
    },
    {
        "story_id": "US-OEM-003",
        "story_name": "Invoice Financing (OEM3)",
        "portal": "OEM Portal",
        "total": 50,
        "catalog": ("json_std", REPORTS / "us-oem-003-cases.json"),
        "report_md": REPORTS / "invoice-financing-regression-execution-report.md",
        "logs": [
            REPORTS / "invoice-financing-slot1-run.log",
            REPORTS / "invoice-financing-slot2-run.log",
        ],
        "dealer": "ved.prakash (OEM Maker)",
        "run_dates": SAT,
    },
    {
        "story_id": "US-DLR-009",
        "story_name": "Purchase Orders & VINs (DLR9)",
        "portal": "Dealer Portal",
        "total": 50,
        "catalog": ("ts", ROOT / "testData/hfi/purchaseOrdersVINsCatalog.ts"),
        "report_md": REPORTS / "purchase-orders-vins-regression-execution-report.md",
        "logs": [
            REPORTS / "purchase-orders-vins-batch1-run.log",
            REPORTS / "purchase-orders-vins-batch2-run.log",
            REPORTS / "purchase-orders-vins-batch3-run.log",
            REPORTS / "purchase-orders-vins-mp00007-run.log",
        ],
        "dealer": "PB00018 (dealer8)",
        "run_dates": SAT,
    },
    {
        "story_id": "US-DLR-010",
        "story_name": "Transaction History (DLR10)",
        "portal": "Dealer Portal",
        "total": 50,
        "catalog": ("ts", ROOT / "testData/hfi/transactionHistoryCatalog.ts"),
        "report_md": REPORTS / "transaction-history-regression-execution-report.md",
        "logs": [
            REPORTS / "transaction-history-regression-dev-run.log",
            REPORTS / "transaction-history-regression-dev-run2.log",
        ],
        "dealer": "PB00018 (dealer8)",
        "run_dates": SAT,
    },
    {
        "story_id": "US-DLR-011",
        "story_name": "Exposure Settings (DLR11)",
        "portal": "Dealer Portal",
        "total": 50,
        "catalog": ("json_std", REPORTS / "us-dlr-011-cases.json"),
        "report_md": REPORTS / "exposure-settings-regression-execution-report.md",
        "logs": [
            REPORTS / "exposure-slot1-run.log",
            REPORTS / "exposure-slot1-rerun.log",
            REPORTS / "exposure-slot1-rerun2.log",
            REPORTS / "exposure-slot1-dp2e-run.log",
            REPORTS / "exposure-slot2-run.log",
            REPORTS / "exposure-slot3-run.log",
            REPORTS / "exposure-slot4-run.log",
            REPORTS / "exposure-slot5-run.log",
            REPORTS / "exposure-priority-batch1-final.log",
            REPORTS / "exposure-priority-batch1-run.log",
            ROOT / "terminals" / "498559.txt",
        ],
        "dealer": "PB00018 (Sat) / dp1056723eaf82453e (Sun)",
        "run_dates": f"{SAT} + {SUN}",
    },
    {
        "story_id": "US-DLR-001",
        "story_name": "Credentials & Authentication (DLR1)",
        "portal": "Dealer Portal",
        "total": 50,
        "catalog": ("ts_dlr1", ROOT / "testData/hfi/dealerLoginCatalog.ts"),
        "report_md": None,
        "logs": [
            REPORTS / "dealer-login-slot1-run.log",
            REPORTS / "dealer-login-slot2-run.log",
            REPORTS / "dealer-login-slot3-run.log",
            REPORTS / "dealer-login-sprint2-run.txt",
        ],
        "dealer": "dp2e166d012ccecc7b (Sun — account locked)",
        "run_dates": SUN,
    },
]

LOG_LINE = re.compile(
    r"^\s*(ok|x|-)\s+\d+.*?(TC-\d{3})",
    re.IGNORECASE,
)
MD_STATUS = re.compile(
    r"\|\s*(TC-\d{3})\s*\|[^|]*\|\s*[^|]*\|\s*\*\*(PASS|FAIL|BLOCKED|SKIPPED|NOT EXECUTED)\*\*",
    re.IGNORECASE,
)
MD_DETAIL = re.compile(
    r"\|\s*(TC-\d{3})\s*\|([^|]+)\|([^|]+)\|\s*\*\*(PASS|FAIL|BLOCKED|SKIPPED)\*\*\s*\|\s*([^|]+)\|([^|]+)\|",
    re.IGNORECASE,
)
ATTACHMENT = re.compile(
    r"(test-results[/\\][^\s]+\.(?:png|webm|md))",
    re.IGNORECASE,
)
STEP_LINE = re.compile(r"^\[STEP\].*$", re.IGNORECASE)
INFO_LINE = re.compile(r"^\[(?:INFO|WARN|ERROR)\].*$", re.IGNORECASE)
SLOT_LINE = re.compile(r"^Slot\s+(\d+)", re.IGNORECASE)
SLOT_FILE = re.compile(r"(?:slot|batch)(\d+)", re.IGNORECASE)
ANSI_ESCAPE = re.compile(r"\x1b\[[0-9;]*m")
ILLEGAL_XLSX = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def sanitize_cell(value) -> str:
    if value is None:
        return ""
    text = str(value)
    text = ANSI_ESCAPE.sub("", text)
    text = ILLEGAL_XLSX.sub("", text)
    return text


def read_log_text(path: Path) -> str:
    raw = path.read_bytes()
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        return raw.decode("utf-16", errors="replace")
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig", errors="replace")
    return raw.decode("utf-8", errors="replace")


def slot_from_log(path: Path, text: str) -> str:
    first = next((ln.strip() for ln in text.splitlines() if ln.strip()), "")
    m = SLOT_LINE.match(first)
    if m:
        return f"Slot {m.group(1)}"
    m = SLOT_FILE.search(path.name)
    if m:
        return f"Slot/Batch {m.group(1)}"
    m = re.search(r"tc(\d{3})-(\d{3})", path.name, re.I)
    if m:
        return f"TC-{m.group(1)}–{m.group(2)}"
    return path.stem


def load_catalog(kind: str, path: Path) -> list[dict]:
    if not path.exists():
        return []
    text = path.read_text(encoding="utf-8", errors="replace")
    if kind == "json_adm":
        data = json.loads(text)
        cases = data.get("cases", data)
        out = []
        for c in cases:
            out.append(
                {
                    "tc_id": c.get("TC ID", c.get("id", "")),
                    "title": c.get("Test Case Title", c.get("title", "")),
                    "type": c.get("Type", c.get("type", "")),
                    "priority": c.get("Priority", c.get("priority", "")),
                    "ac": c.get("AC / Flow Ref", c.get("ac", "")),
                    "preconditions": c.get("Preconditions", c.get("preconditions", "")),
                    "steps": c.get("Test Steps", c.get("steps", "")),
                    "test_data": c.get("Test Data", c.get("testData", "-")),
                    "expected": c.get("Expected Result", c.get("expected", c.get("expectedResult", ""))),
                }
            )
        return out
    if kind == "json_oem2":
        cases = json.loads(text)
        out = []
        for c in cases:
            out.append(
                {
                    "tc_id": c.get("TC ID", ""),
                    "title": c.get("Test Case Title", ""),
                    "type": c.get("Type", ""),
                    "priority": c.get("Priority", ""),
                    "ac": c.get("AC / Flow Ref", ""),
                    "preconditions": c.get("Preconditions", ""),
                    "steps": c.get("Test Steps", ""),
                    "test_data": c.get("Test Data", "-"),
                    "expected": c.get("Expected Result", ""),
                }
            )
        return out
    if kind == "json_std":
        cases = json.loads(text)
        return [
            {
                "tc_id": c.get("id", ""),
                "title": c.get("title", ""),
                "type": c.get("type", ""),
                "priority": c.get("priority", ""),
                "ac": c.get("ac", ""),
                "preconditions": c.get("preconditions", ""),
                "steps": c.get("steps", ""),
                "test_data": "-",
                "expected": c.get("expected", ""),
            }
            for c in cases
        ]
    if kind == "ts_dlr1":
        blocks = re.findall(
            r'\{\s*"id":\s*"(TC-\d{3})".*?"title":\s*"([^"]+)".*?"type":\s*"([^"]+)".*?"priority":\s*"([^"]+)".*?"acFlowRef":\s*"([^"]*)".*?"preconditions":\s*"([^"]*)".*?"steps":\s*"([^"]*)".*?"testData":\s*"([^"]*)".*?"expectedResult":\s*"([^"]*)"',
            text,
            re.DOTALL,
        )
        return [
            {
                "tc_id": b[0],
                "title": b[1],
                "type": b[2],
                "priority": b[3],
                "ac": b[4],
                "preconditions": b[5].replace("\\n", "\n"),
                "steps": b[6].replace("\\n", "\n"),
                "test_data": b[7],
                "expected": b[8].replace("\\n", "\n"),
            }
            for b in blocks
        ]
    if kind == "ts":
        rows = re.findall(
            r'\{\s*id:\s*"(TC-\d{3})",\s*title:\s*"([^"]+)",\s*type:\s*"([^"]+)",\s*priority:\s*"([^"]+)",\s*ac:\s*"([^"]*)"\s*\}',
            text,
        )
        return [
            {
                "tc_id": r[0],
                "title": r[1],
                "type": r[2],
                "priority": r[3],
                "ac": r[4],
                "preconditions": "",
                "steps": "",
                "test_data": "-",
                "expected": "",
            }
            for r in rows
        ]
    return []


def parse_logs(log_paths: list[Path]) -> dict[str, dict]:
    results: dict[str, dict] = {}
    ts_re = re.compile(r"\[(\d{4}-\d{2}-\d{2}T[\d:.]+Z)\]")

    def entry_ts(entry: dict, fallback: float) -> float:
        for key in ("execution_steps", "log_line"):
            for m in ts_re.finditer(entry.get(key, "")):
                try:
                    return datetime.fromisoformat(m.group(1).replace("Z", "+00:00")).timestamp()
                except ValueError:
                    continue
        return fallback

    for lp in log_paths:
        if not lp.exists():
            continue
        text = read_log_text(lp)
        run_date = SAT if SAT in text[:500] or "Aug 15" in text[:800] else ""
        if SUN in text or "Aug 16" in text or "2026-08-16" in text:
            run_date = SUN if not run_date else f"{run_date}, {SUN}"
        file_mtime = lp.stat().st_mtime
        if not run_date:
            run_date = datetime.fromtimestamp(file_mtime).strftime("%Y-%m-%d")
        batch_slot = slot_from_log(lp, text)
        pending_steps: list[str] = []
        lines = text.splitlines()
        for idx, line in enumerate(lines):
            stripped = line.strip()
            if STEP_LINE.match(stripped) or INFO_LINE.match(stripped):
                pending_steps.append(stripped[:500])
                continue
            m = LOG_LINE.match(line)
            if not m:
                continue
            sym, tc = m.group(1).lower(), m.group(2).upper()
            status = {"ok": "PASS", "x": "FAIL", "-": "BLOCKED/SKIPPED"}.get(sym, "UNKNOWN")
            error_lines: list[str] = []
            if sym == "x":
                for follow in lines[idx + 1 : idx + 25]:
                    fs = follow.strip()
                    if LOG_LINE.match(follow):
                        break
                    if fs and (
                        fs.startswith("Error")
                        or "expect(" in fs
                        or fs.startswith("Timeout")
                        or "Received:" in fs
                        or fs.startswith("Call log:")
                    ):
                        error_lines.append(fs[:400])
            chunk = "\n".join(lines[max(0, idx - 40) : idx + 30])
            shots = ATTACHMENT.findall(chunk)
            step_log = "\n".join(pending_steps[-20:])
            if error_lines:
                step_log = (step_log + "\n" if step_log else "") + "\n".join(error_lines[:8])
            candidate = {
                "status": status,
                "log_file": str(lp.relative_to(ROOT)),
                "run_date": run_date,
                "log_line": stripped[:300],
                "slot": batch_slot,
                "execution_steps": step_log[:4000],
                "screenshots": "; ".join(dict.fromkeys(shots)),
            }
            prev = results.get(tc)
            if not prev or entry_ts(candidate, file_mtime) >= entry_ts(prev, 0):
                results[tc] = candidate
            pending_steps = []
    return results


def parse_markdown_report(path: Path | None) -> dict[str, dict]:
    if not path or not path.exists():
        return {}
    text = path.read_text(encoding="utf-8", errors="replace")
    out: dict[str, dict] = {}
    for m in MD_DETAIL.finditer(text):
        tc, _title, _classif, status, actual, root = m.groups()
        out[tc.upper()] = {
            "status": status.upper(),
            "actual_result": actual.strip(),
            "root_cause": root.strip(),
            "source": str(path.relative_to(ROOT)),
        }
    for m in MD_STATUS.finditer(text):
        tc, status = m.groups()
        tc = tc.upper()
        if tc not in out:
            out[tc] = {"status": status.upper(), "source": str(path.relative_to(ROOT))}
    return out


def find_pass_evidence(story_id: str, tc_id: str) -> str:
    file_name = f"{story_id}-{tc_id}-PASS.png"
    path = SPRINT2_EVIDENCE_DIR / file_name
    if path.exists():
        return str(path.relative_to(ROOT)).replace("\\", "/")
    return ""


def find_evidence(tc_id: str, title: str, story_id: str = "") -> str:
    if story_id:
        pass_shot = find_pass_evidence(story_id, tc_id)
        if pass_shot:
            return pass_shot
    if not TEST_RESULTS.exists():
        return ""
    slug = re.sub(r"[^a-z0-9]+", "-", title.lower())[:40]
    tc_num = tc_id.replace("TC-", "")
    hits = []
    for p in TEST_RESULTS.rglob("*"):
        if not p.is_file():
            continue
        if p.suffix.lower() not in {".png", ".webm", ".md"}:
            continue
        name = p.name.lower()
        parent = str(p.parent).lower()
        if f"tc-{tc_num}" in parent or f"tc-{tc_num}" in name or slug[:20] in parent:
            hits.append(str(p.relative_to(ROOT)))
    return "; ".join(hits[:5])


def style_header(ws, row=1):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF", size=11)
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def autosize(ws, max_width=60):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min(max(len(str(c.value or "")) for c in col) + 2, max_width)
        ws.column_dimensions[letter].width = width


def set_file_hyperlink(cell, rel_path: str) -> None:
    rel_path = (rel_path or "").strip()
    if not rel_path:
        return
    first = rel_path.split(";")[0].strip()
    target = (ROOT / first).resolve()
    cell.value = first
    if not target.exists():
        return
    excel_dir = OUT_XLSX.parent.resolve()
    try:
        link_target = target.relative_to(excel_dir).as_posix()
    except ValueError:
        link_target = target.as_uri()
    cell.hyperlink = link_target
    cell.font = Font(color="0563C1", underline="single")


def resolve_suite_statuses(suite: dict) -> dict[str, str]:
    log_res = parse_logs(suite["logs"])
    md_res = parse_markdown_report(suite["report_md"])
    statuses: dict[str, str] = {}
    for i in range(1, suite["total"] + 1):
        tc_id = f"TC-{i:03d}"
        lr = log_res.get(tc_id, {})
        mr = md_res.get(tc_id, {})
        log_status = lr.get("status", "")
        md_status = mr.get("status", "")
        if log_status in ("PASS", "FAIL", "BLOCKED/SKIPPED"):
            status = log_status
        elif md_status and md_status not in ("SKIPPED", "NOT EXECUTED"):
            status = md_status
        elif md_status == "SKIPPED":
            status = "BLOCKED/SKIPPED"
        elif log_status:
            status = log_status
        else:
            status = "NOT EXECUTED"
        if status in ("BLOCKED", "SKIPPED"):
            status = "BLOCKED/SKIPPED"
        statuses[tc_id] = status
    return statuses


def export_pass_manifest() -> dict[str, list[str]]:
    manifest: dict[str, list[str]] = {}
    for suite in SUITES:
        if suite["story_id"] == "UC-ADM-001":
            continue
        statuses = resolve_suite_statuses(suite)
        manifest[suite["story_id"]] = sorted(
            tc for tc, status in statuses.items() if status == "PASS"
        )
    REPORTS.mkdir(parents=True, exist_ok=True)
    PASS_MANIFEST.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return manifest


def main():
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Executive Summary"

    headers_detail = [
        "Sprint",
        "User Story ID",
        "User Story Name",
        "Portal",
        "TC ID",
        "Test Case Title",
        "Type",
        "Priority",
        "AC / Flow Ref",
        "Preconditions",
        "Test Steps",
        "Test Data",
        "Expected Result",
        "Execution Date(s)",
        "Slot / Batch",
        "Status",
        "Actual Result",
        "Execution Step Log",
        "Root Cause / Remarks",
        "Defect Type",
        "Account / Dealer Used",
        "Environment",
        "Log File",
        "Log Evidence Line",
        "Screenshot Path(s)",
        "Video / Trace Path",
        "Ortoni Report",
        "Markdown Report",
    ]
    ws_detail = wb.create_sheet("All Test Cases")
    ws_detail.append(headers_detail)
    style_header(ws_detail)

    sum_headers = [
        "User Story ID",
        "User Story Name",
        "Total TCs",
        "PASS",
        "FAIL",
        "BLOCKED/SKIPPED",
        "NOT EXECUTED",
        "Pass %",
        "Run Window",
        "Primary Log(s)",
    ]
    ws_sum.append(sum_headers)
    style_header(ws_sum)

    all_rows = []
    detail_screenshot_links: list[tuple[int, str]] = []
    story_sheet_links: dict[str, list[tuple[int, str]]] = {}
    grand = {"total": 0, "pass": 0, "fail": 0, "blocked": 0, "ne": 0}
    SCREENSHOT_COL = 25

    for suite in SUITES:
        catalog = load_catalog(suite["catalog"][0], suite["catalog"][1])
        log_res = parse_logs(suite["logs"])
        md_res = parse_markdown_report(suite["report_md"])
        counts = {"pass": 0, "fail": 0, "blocked": 0, "ne": 0}

        for i in range(1, suite["total"] + 1):
            tc_id = f"TC-{i:03d}"
            cat = next((c for c in catalog if c["tc_id"] == tc_id), None)
            title = cat["title"] if cat else f"(see Excel source — {tc_id})"
            lr = log_res.get(tc_id, {})
            mr = md_res.get(tc_id, {})

            # Prefer live log outcome; markdown supplements curated analysis
            log_status = lr.get("status", "")
            md_status = mr.get("status", "")
            if log_status in ("PASS", "FAIL", "BLOCKED/SKIPPED"):
                status = log_status
            elif md_status and md_status not in ("SKIPPED", "NOT EXECUTED"):
                status = md_status
            elif md_status == "SKIPPED":
                status = "BLOCKED/SKIPPED"
            elif log_status:
                status = log_status
            else:
                status = "NOT EXECUTED"

            if status == "PASS":
                counts["pass"] += 1
            elif status == "FAIL":
                counts["fail"] += 1
            elif status in ("BLOCKED", "BLOCKED/SKIPPED", "SKIPPED"):
                counts["blocked"] += 1
                status = "BLOCKED/SKIPPED"
            else:
                counts["ne"] += 1

            pass_shot = find_pass_evidence(suite["story_id"], tc_id)
            shots = pass_shot or lr.get("screenshots", "") or find_evidence(
                tc_id, title, suite["story_id"]
            )
            videos = "; ".join([s for s in shots.split("; ") if s.endswith(".webm")])
            pngs = "; ".join([s for s in shots.split("; ") if s.endswith(".png")])
            if status == "PASS" and pass_shot:
                pngs = pass_shot

            defect_type = ""
            if status == "FAIL":
                rc = (mr.get("root_cause") or "").lower()
                if "automation" in rc:
                    defect_type = "Automation Issue"
                elif "application" in rc or "defect" in rc:
                    defect_type = "Application Defect"
                elif "integration" in rc:
                    defect_type = "Integration Defect"
                else:
                    defect_type = "See Root Cause"
            elif status == "BLOCKED/SKIPPED":
                defect_type = "Dependency / Test Data"

            slot_label = lr.get("slot") or (
                mr.get("source", "").split("/")[-1].replace(".md", "") if mr else "See report"
            )
            actual = mr.get("actual_result", "") or (
                lr.get("log_line", "") if status in ("PASS", "FAIL") else ""
            )
            root_cause = mr.get("root_cause", "") or (
                lr.get("execution_steps", "")[:500] if status == "FAIL" else lr.get("log_line", "")
            )
            exec_log = lr.get("execution_steps", "")

            row = [
                sanitize_cell(v) for v in [
                "Sprint 2",
                suite["story_id"],
                suite["story_name"],
                suite["portal"],
                tc_id,
                title,
                cat.get("type", "") if cat else "",
                cat.get("priority", "") if cat else "",
                cat.get("ac", "") if cat else "",
                cat.get("preconditions", "") if cat else "",
                cat.get("steps", "") if cat else "",
                cat.get("test_data", "-") if cat else "-",
                cat.get("expected", "") if cat else "",
                lr.get("run_date") or suite["run_dates"],
                slot_label,
                status,
                actual,
                exec_log,
                root_cause,
                defect_type,
                suite["dealer"],
                "DEV — http://devdealerportal.centralindia.cloudapp.azure.com",
                lr.get("log_file", mr.get("source", "")),
                lr.get("log_line", ""),
                pngs,
                videos,
                "ortoni-report/ortoni-report.html",
                str(suite["report_md"].relative_to(ROOT)) if suite["report_md"] else "",
            ]]
            ws_detail.append(row)
            if pngs:
                detail_screenshot_links.append((ws_detail.max_row, pngs))
            all_rows.append(row)

        total = suite["total"]
        pct = round(counts["pass"] / total * 100, 1) if total else 0
        ws_sum.append(
            [
                suite["story_id"],
                suite["story_name"],
                total,
                counts["pass"],
                counts["fail"],
                counts["blocked"],
                counts["ne"],
                pct,
                suite["run_dates"],
                "; ".join(p.name for p in suite["logs"][:3] if p.exists()),
            ]
        )
        grand["total"] += total
        grand["pass"] += counts["pass"]
        grand["fail"] += counts["fail"]
        grand["blocked"] += counts["blocked"]
        grand["ne"] += counts["ne"]

    ws_sum.append([])
    ws_sum.append(
        [
            "GRAND TOTAL",
            "All Sprint 2 Stories",
            grand["total"],
            grand["pass"],
            grand["fail"],
            grand["blocked"],
            grand["ne"],
            round(grand["pass"] / grand["total"] * 100, 1) if grand["total"] else 0,
            f"{SAT} – {SUN}",
            "See Detail sheet",
        ]
    )

    for row_idx, shot_path in detail_screenshot_links:
        set_file_hyperlink(ws_detail.cell(row=row_idx, column=SCREENSHOT_COL), shot_path)

    # Per-story sheets (abbreviated columns)
    for suite in SUITES:
        safe = re.sub(r"[^\w]", "_", suite["story_id"])[:31]
        ws = wb.create_sheet(safe)
        story_sheet_links[safe] = []
        short_h = [
            "TC ID",
            "Title",
            "Status",
            "Slot/Batch",
            "Actual Result",
            "Execution Steps",
            "Root Cause",
            "Log",
            "Screenshots",
        ]
        ws.append(short_h)
        style_header(ws)
        for row in all_rows:
            if row[1] == suite["story_id"]:
                ws.append(
                    [row[4], row[5], row[15], row[14], row[16], row[17], row[18], row[22], row[24]]
                )
                if row[24]:
                    story_sheet_links[safe].append((ws.max_row, row[24]))
        for row_idx, shot_path in story_sheet_links.get(safe, []):
            set_file_hyperlink(ws.cell(row=row_idx, column=9), shot_path)
        autosize(ws)

    autosize(ws_detail, 45)
    autosize(ws_sum, 30)

    meta = wb.create_sheet("README")
    meta.append(["Sprint 2 Regression Management Report"])
    meta.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    meta.append(["Run window", f"Saturday {SAT} + Sunday {SUN}"])
    meta.append(["Total test cases", grand["total"]])
    meta.append([])
    meta.append(["User stories included"])
    for s in SUITES:
        meta.append([s["story_id"], s["story_name"], s["total"], s["run_dates"]])
    meta.append([])
    meta.append(["Evidence locations"])
    meta.append(["Ortoni HTML", "ortoni-report/ortoni-report.html"])
    meta.append(["Playwright HTML", "my-report/index.html"])
    meta.append(["PASS evidence screenshots", "reports/sprint2-evidence/{STORY}-TC-NNN-PASS.png"])
    meta.append(["PASS manifest", str(PASS_MANIFEST.relative_to(ROOT))])
    meta.append(["Text logs", "reports/*.log"])
    meta.append([])
    meta.append(["Note", "DLR-001 (Sun) account was LOCKED — many login TCs show FAIL/NOT EXECUTED."])
    meta.append(["Note", "DLR-011 Sun run used dealer dp1056723eaf82453e; Sat used PB00018."])

    REPORTS.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    print(f"Wrote {OUT_XLSX}")
    print(f"Total TCs: {grand['total']} | PASS: {grand['pass']} | FAIL: {grand['fail']} | BLOCKED: {grand['blocked']} | NE: {grand['ne']}")


if __name__ == "__main__":
    import sys

    if "--export-pass-manifest" in sys.argv:
        manifest = export_pass_manifest()
        total = sum(len(v) for v in manifest.values())
        print(f"Wrote {PASS_MANIFEST} ({total} PASS TCs across 6 stories)")
    else:
        export_pass_manifest()
        main()
