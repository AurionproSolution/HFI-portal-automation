#!/usr/bin/env python3
"""Apply Sprint 3 test case enhancements to the 149-TC Excel workbook."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

XLSX = Path(r"c:\Users\piyush.more\Downloads\Honda_Finance_Sprint3_Test_Cases_Updated_149.xlsx")
FALLBACK_XLSX = Path(__file__).resolve().parents[1] / "reports" / "Honda_Finance_Sprint3_Test_Cases_Updated_149.xlsx"

STORY_SHEETS = [
    "US-DLR-012",
    "US-DLR-013",
    "US-DLR-014",
    "US-DLR-015",
    "US-DLR-016",
    "UC-OEM-004",
    "US-OPS-002",
    "US-OPS-003",
]

COL = {
    "module": 1,
    "id": 2,
    "title": 3,
    "type": 4,
    "priority": 5,
    "ac": 6,
    "preconditions": 7,
    "steps": 8,
    "test_data": 9,
    "expected": 10,
    "status": 11,
    "remarks": 12,
}

SYSTEM_DOWN_MSG = (
    "We're having trouble loading your data. We can't connect to one or more of the "
    "systems that provide this information. Please try again in a few minutes."
)

UPDATES: dict[str, dict] = {
    "TC_CFD_002": {
        "test_data": "Dealer with Consumer Finance access; target and MTD disbursal data available in CFS",
        "expected": (
            "1. Disbursed (MTD) card displays MTD amount against the Honda-set target.\n"
            "2. Achievement percentage is calculated per defined business rule (Disbursed (MTD) ÷ Target).\n"
            "3. Residual amount is calculated per defined business rule (Target − Disbursed (MTD)).\n"
            "4. Dealer Payout (MTD) card displays commissions earned in the current cycle."
        ),
    },
    "TC_CFD_003": {
        "test_data": "Dealer scope where disbursed amount meets or exceeds configured target",
        "expected": (
            "1. Achievement percentage shows 100% or greater when target is met or exceeded.\n"
            "2. Residual amount displays as nil rather than a negative value."
        ),
    },
    "TC_CFD_004": {
        "test_data": "Dealer scope with current-period and same-date-last-month baseline data available",
        "expected": (
            "1. Same-date-last-month baseline comparison displays for the selected scope.\n"
            "2. Variance indicators display per configured baseline data."
        ),
    },
    "TC_APP_014": {
        "title": "Verify empty state when search/scope returns zero records and when section or detail panel data is unavailable",
        "ac": "AC-11, AC-25, E1",
        "steps": (
            "1. Apply a search term or scope that returns zero ledger records.\n"
            "2. Observe the ledger empty-state behaviour.\n"
            "3. Open Application Detail for a case where one panel/section has no data.\n"
            "4. Observe the affected section or panel."
        ),
        "test_data": "Search term with no matches; application with missing section data",
        "expected": (
            "1. Ledger displays the defined empty state when no matching records exist.\n"
            "2. For missing section or panel data (AC-25), the affected section/panel is not displayed "
            "or shows an empty state while remaining sections continue to render."
        ),
        "remarks": "Empty State & Partial Nil Sections",
    },
    "TC_APP_020": {
        "test_data": "Applications screen with KPI cards and ledger containing monetary values",
        "expected": (
            "Currency and numbering values across KPI cards and the ledger display in Indian numbering "
            "format per defined business rules."
        ),
    },
    "TC_DIS_001": {
        "test_data": "Dealer with disbursal data for selected branch scope and period",
        "expected": (
            "1. Disbursed (MTD), Total Value, Cycle Time and Avg Ticket Size KPI cards are displayed.\n"
            "2. Cycle Time is expressed in days to one decimal place.\n"
            "3. Avg Ticket Size equals Total Value divided by the number of disbursed cases in scope."
        ),
    },
    "TC_DIS_002": {
        "test_data": "Dealer scope with current and previous period analytics data",
        "expected": (
            "1. Current and Previous tiles show disbursed value for the selected period and the immediately "
            "preceding period of equal length.\n"
            "2. Delta shows absolute variance with direction indicator and percentage variance calculated on the Previous value."
        ),
    },
    "TC_DIS_004": {
        "test_data": "Dealer scope with week-wise disbursal analytics data",
        "expected": (
            "Peak disbursal bucket is highlighted and percentage variance displays per configured analytics data."
        ),
    },
    "TC_DIS_017": {
        "test_data": "Disbursements screen with KPI cards, analytics and ledger monetary values",
        "expected": (
            "Disbursal amounts display in Indian numbering format with appropriate decimal precision per business rules."
        ),
    },
    "TC_PAY_009": {
        "ac": "AC-22, AC-26, A2",
        "steps": (
            "1. Upload and parse a valid payout file.\n"
            "2. Discard the parsed file before submission.\n"
            "3. Separately, observe a submitted batch that is rejected downstream.\n"
            "4. Verify dealer notification and ledger status."
        ),
        "expected": (
            "1. Discarding parsed file returns upload panel to empty state without storing payout values.\n"
            "2. When a submitted batch is rejected downstream (AC-26), the dealer is notified and can correct "
            "and re-upload; invoices are marked Rejected in the ledger."
        ),
        "remarks": "Discard Upload & Rejection Lifecycle",
    },
    "TC_PAY_019": {
        "ac": "AC-23, AC-24",
        "steps": (
            "1. Upload a payout file and observe parsing loading state.\n"
            "2. With a payout batch already submitted and pending approval, attempt to upload a new file.\n"
            "3. Observe system behaviour."
        ),
        "expected": (
            "1. Loading state is displayed during parsing and submit actions are disabled.\n"
            "2. When a batch is pending approval (AC-24), new upload is blocked or clearly distinguished "
            "from the pending batch per defined concurrency rule."
        ),
        "remarks": "Loading State & Concurrent Upload",
    },
    "TC_PAY_020": {
        "ac": "AC-25, AC-27, AC-28",
        "steps": (
            "1. Inspect Payout Ledger columns and status display.\n"
            "2. Verify deduction formatting.\n"
            "3. After downstream approval of a submitted batch, verify ledger and Payout Amount Status."
        ),
        "expected": (
            "1. Ledger displays required columns with correct status display and deduction formatting.\n"
            "2. When a submitted batch is approved downstream (AC-25), entries appear as Processed in the "
            "Payout Ledger and contribute to Total Approved in Payout Amount Status."
        ),
        "remarks": "Payout Ledger & Downstream Approval",
    },
    "TC_PAY_023": {
        "ac": "AC-30, AC-31, E6",
        "expected": (
            "1. When submission cannot be saved (AC-30), no partial state is written and no payout values are stored.\n"
            "2. Dealer is prompted to retry.\n"
            "3. Failure is communicated clearly when file store service is unavailable."
        ),
        "remarks": "Atomic Failure & File Store Error",
    },
    "TC_CHK_001": {
        "title": "Verify Checker Console page layout and session context upon opening Checker Queue",
        "ac": "AC-2, Main Flow Step 1",
        "steps": (
            "1. Sign in as an authenticated OEM Checker.\n"
            "2. Open the Checker Queue.\n"
            "3. Verify Checker Console tab, breadcrumb, page title and description.\n"
            "4. Verify signed-in user name and role are displayed."
        ),
        "test_data": "Valid OEM Checker login credentials",
        "expected": (
            "1. Checker Console tab is active.\n"
            "2. Breadcrumb, page title and description are displayed.\n"
            "3. Signed-in user's name and role are displayed."
        ),
        "remarks": "Checker Console UI",
    },
    "TC_CHK_002": {
        "ac": "AC-3, Main Flow Step 2",
        "preconditions": "One or more financing requests are awaiting Checker approval.",
        "test_data": "Checker queue with pending approval requests",
        "expected": (
            "Pending Review, Total Amount and Distinct Dealers KPI cards are displayed and reflect "
            "the requests currently awaiting approval."
        ),
        "remarks": "Checker KPIs",
    },
    "TC_CHK_003": {
        "ac": "AC-5, Main Flow Step 3",
        "steps": (
            "1. Inspect the Awaiting Your Approval list.\n"
            "2. Verify each row shows selection checkbox, Invoice/Batch (batch ID beneath invoice number), "
            "Dealer (dealer code beneath dealer name), Product (category tag with View Details), Amount, "
            "Available Limit, Maker submission details and Actions (Approve/Reject, Approve Selected).\n"
            "3. Verify Select All control is available."
        ),
        "expected": (
            "All required columns and controls are displayed per AC-5 and main-flow UI specification. "
            "Available limit reflects LMS data for the dealer."
        ),
        "remarks": "Approval Queue Grid",
    },
    "TC_CHK_004": {
        "ac": "Main Flow Step 4",
        "steps": (
            "1. Select View Details on a row's Product column.\n"
            "2. Inspect the product breakdown displayed."
        ),
        "test_data": "Pending request with multi-model product breakdown",
        "expected": (
            "Full model/quantity breakdown for the selected request is displayed so the Checker can "
            "inspect underlying units before deciding, without leaving the queue."
        ),
        "remarks": "Product Details Modal",
    },
    "TC_CHK_005": {
        "title": "Verify single request approval, LMS disbursement submission, loading state and Maker notification",
        "ac": "AC-6, AC-11, AC-14, Main Flow Step 5, Step 7",
        "steps": (
            "1. Click Approve on a single pending request within available dealer limit.\n"
            "2. Observe loading state while submission is in progress.\n"
            "3. Confirm Approve/Reject actions are disabled during processing.\n"
            "4. Verify queue update and Maker notification after completion."
        ),
        "test_data": "Single pending request within dealer available limit (Mock LMS APIs for Sprint 3)",
        "expected": (
            "1. Request is submitted to LMS for disbursement (Mock APIs used for Sprint 3).\n"
            "2. Loading state is displayed and Approve/Reject actions are disabled during submission (AC-14).\n"
            "3. Item is removed from awaiting-approval list and KPI cards recalculate (AC-11).\n"
            "4. Maker is notified with approval outcome and LAN details where returned (Main Flow Step 7)."
        ),
        "remarks": "Single Approval Flow",
    },
    "TC_CHK_006": {
        "ac": "AC-7, Main Flow Step 5",
        "steps": (
            "1. Select multiple requests individually or via Select All.\n"
            "2. Click Approve Selected.\n"
            "3. Confirm the action."
        ),
        "test_data": "Multiple selected pending requests (Mock LMS APIs for Sprint 3)",
        "expected": (
            "All selected requests are submitted to LMS for disbursement (Mock APIs used for Sprint 3), "
            "removed from the queue and Maker is notified per approval outcome."
        ),
        "remarks": "Bulk Approval Flow",
    },
    "TC_CHK_007": {
        "ac": "AC-5, A2",
        "expected": (
            "1. Select All checks all visible row checkboxes.\n"
            "2. Deselecting specific items before approval leaves only the intended items selected (Alt Flow A2)."
        ),
        "remarks": "Select All Logic",
    },
    "TC_CHK_008": {
        "ac": "AC-8, AC-9, E1, Main Flow Step 6",
        "steps": (
            "1. Click Reject on a pending request.\n"
            "2. Attempt to submit with blank remarks.\n"
            "3. Enter mandatory remarks and submit."
        ),
        "test_data": "Pending request in Checker queue",
        "expected": (
            "1. Rejection is blocked until remarks are entered (AC-8, E1).\n"
            "2. With remarks entered, rejection is recorded, item is removed from queue, KPIs recalculate "
            "and Maker is notified via email/SMS with rejection remarks (AC-9)."
        ),
        "remarks": "Rejection Flow & Remarks",
    },
    "TC_CHK_009": {
        "title": "Verify Audit and Tracking log displays latest 50 Checker actions on screen",
        "type": "UI / Functional",
        "ac": "AC-20, Business Rules",
        "preconditions": "Checker has performed one or more actions in the portal.",
        "steps": (
            "1. Perform an approve or reject action in Checker Queue.\n"
            "2. Navigate to Audit and Tracking.\n"
            "3. Inspect the action log displayed on screen."
        ),
        "test_data": "Checker session with recent approval/rejection actions",
        "expected": (
            "Audit and Tracking log displays the latest 50 Checker actions on screen, including recent "
            "approval/rejection decisions."
        ),
        "remarks": "Audit Log UI",
    },
    "TC_CHK_010": {
        "title": "Verify access denial for non-Checker users and segregation of duties for OEM Maker",
        "ac": "AC-1, Business Rules",
        "steps": (
            "Part A – Non-Checker access (AC-1):\n"
            "1. Sign in as a user without the Checker role.\n"
            "2. Attempt to access the Checker Queue.\n\n"
            "Part B – Maker segregation (Business Rules):\n"
            "3. Sign in as OEM Maker.\n"
            "4. Attempt to navigate to Checker Queue."
        ),
        "test_data": "Non-Checker user credentials; OEM Maker credentials",
        "expected": (
            "Part A: Access is denied and message 'Account does not exist' is displayed.\n"
            "Part B: OEM Maker cannot access Checker Queue (segregation of duties enforced)."
        ),
        "remarks": "Role Access Control",
    },
    "TC_CHK_011": {
        "ac": "AC-13, A3",
        "steps": (
            "1. Search queue by invoice number.\n"
            "2. Search by dealer.\n"
            "3. Search by Maker.\n"
            "4. Verify filtered results."
        ),
        "test_data": "Populated Checker queue; partial search terms",
        "expected": (
            "Queue filters to matching requests only for each search field (invoice, dealer, Maker) per AC-13 and Alt Flow A3."
        ),
        "remarks": "Queue Search",
    },
    "TC_CHK_012": {
        "ac": "AC-12, E2, Business Rules",
        "preconditions": "Pending approval requests consume the dealer's available limit.",
        "steps": (
            "1. Observe Available Limit for a dealer whose limit is fully consumed by pending requests.\n"
            "2. As OEM Maker, attempt to upload a further invoice for the same dealer."
        ),
        "test_data": "Dealer with pending requests consuming full available limit",
        "expected": (
            "System displays: 'The invoice(s) amount exceeds the available dealer limit.' "
            "Further invoice upload for that dealer is blocked (AC-12, E2)."
        ),
        "remarks": "Limit Consumption Check",
    },
    "TC_CHK_013": {
        "ac": "AC-4, AC-18, E3",
        "steps": (
            "1. Open Checker Queue when no requests are awaiting approval.\n"
            "2. Observe KPI cards and awaiting-approval list.\n"
            "3. Observe any section where data does not exist."
        ),
        "test_data": "Empty awaiting-approval queue",
        "expected": (
            "1. KPI cards show zero values and awaiting-approval list shows empty state (AC-4, E3).\n"
            "2. Where section data does not exist (AC-18), section shows empty state with panel frame "
            "and KPI labels retained."
        ),
        "remarks": "Empty Queue & Nil Sections",
    },
    "TC_CHK_014": {
        "ac": "AC-16, E5",
        "expected": (
            "Approval or rejection completes successfully even when Maker notification service fails. "
            "Notification failure is logged without blocking the Checker workflow (AC-16, E5)."
        ),
        "remarks": "Notification Failure Handling",
    },
    "TC_CHK_015": {
        "ac": "AC-17, E6",
        "expected": (
            f"Validation message is displayed: \"{SYSTEM_DOWN_MSG}\" "
            "Approve/Reject actions are not offered when queue source is unavailable."
        ),
        "remarks": "System Down Handling",
    },
    "TC_CHK_016": {
        "title": "Verify branch-scoped awaiting-approval list for Checker assigned to specific branches",
        "ac": "A4, Business Rules",
        "preconditions": "Checker is assigned to a defined branch scope; requests exist across multiple branches.",
        "steps": (
            "1. Sign in as branch-scoped OEM Checker.\n"
            "2. Open Checker Queue.\n"
            "3. Review pending requests listed."
        ),
        "test_data": "Checker assigned to specific branch scope",
        "expected": (
            "Awaiting-approval list is limited to the Checker's branch scope (Alt Flow A4). "
            "Requests outside assigned branches are not listed."
        ),
        "remarks": "Branch Scoping (Alt Flow A4)",
    },
    "TC_CHK_017": {
        "ac": "A1, Business Rules",
        "expected": (
            "Checker can approve one item and reject another independently within the same batch. "
            "Each item is actioned separately per partial approval business rule (Alt Flow A1)."
        ),
        "remarks": "Partial Batch Decision",
    },
    "TC_CHK_018": {
        "ac": "AC-15, E4",
        "expected": (
            "When approval/rejection cannot be saved (AC-15, E4), no partial state is written. "
            "Request remains in Pending Review state and Checker must retry."
        ),
        "remarks": "Transaction Rollback",
    },
    "TC_CHK_019": {
        "ac": "AC-11, Main Flow Step 8",
        "preconditions": "Multiple requests awaiting approval in Checker queue.",
        "test_data": "Checker queue with known pending count and total amount",
        "expected": (
            "Actioned item is removed from awaiting-approval list. Pending Review, Total Amount and "
            "Distinct Dealers KPI cards recalculate accordingly (AC-11, Main Flow Step 8)."
        ),
        "remarks": "KPI Auto-Recalculation",
    },
    "TC_CHK_020": {
        "title": "Verify OEM Checker dashboard displays same details as OEM Maker dashboard",
        "type": "UI / Functional",
        "ac": "AC-19",
        "preconditions": "Authenticated OEM Checker user.",
        "steps": (
            "1. Sign in as OEM Checker and view the dashboard.\n"
            "2. Compare dashboard widgets, layout and data scope with OEM Maker dashboard (refer HFIT-186).\n"
            "3. Verify applicable Maker dashboard acceptance criteria apply to Checker dashboard."
        ),
        "test_data": "OEM Checker and OEM Maker sessions for comparison",
        "expected": (
            "Checker dashboard displays the same details as the OEM Maker dashboard. "
            "Applicable HFIT-186 Maker dashboard acceptance criteria are met for the Checker view."
        ),
        "remarks": "Dashboard Parity (AC-19)",
    },
    "TC_CHK_021": {
        "test_data": "Checker queue with monetary values in Amount and Available Limit columns",
        "expected": (
            "Request amounts and available limits display in Indian numbering format with ₹ prefix per business rules."
        ),
        "remarks": "Format Compliance",
    },
    "TC_CHK_022": {
        "ac": "AC-10, Business Rules",
        "steps": (
            "1. Verify state of a rejected request.\n"
            "2. Confirm rejected record cannot be edited or resubmitted in place from Checker UI.\n"
            "3. Confirm Maker must amend and re-upload."
        ),
        "test_data": "Rejected request in system",
        "expected": (
            "Rejected request is not resubmitted in place (AC-10). Maker must amend and re-upload the file as a new record."
        ),
        "remarks": "Rejection Lifecycle",
    },
    "TC_OPS2_001": {
        "title": "Verify access control and page header for Offline Payment Verification",
        "ac": "AC-1, AC-2",
        "steps": (
            "Part A – Access denial (AC-1):\n"
            "1. Sign in as a user without offline payment verification entitlement.\n"
            "2. Attempt to access Ops Console – Offline Payment Verification.\n\n"
            "Part B – Authorized access (AC-2):\n"
            "3. Sign in as authorised Operations user with verification entitlement.\n"
            "4. Open Offline Payment Verification.\n"
            "5. Verify page header elements."
        ),
        "test_data": "Unauthorized Ops user; authorised Operations user with verification entitlement",
        "expected": (
            "Part A: Access is denied and module is not displayed.\n"
            "Part B: Ops Console badge, breadcrumb, page title and description, branch scope chip, "
            "and signed-in user name and role are displayed."
        ),
        "remarks": "Access & Header",
    },
    "TC_OPS2_002": {
        "ac": "AC-3, AC-4, AC-22",
        "steps": (
            "1. Open Offline Payment Verification with data available.\n"
            "2. Verify KPI cards against Pending Verification tab.\n"
            "3. Simulate backend unavailability and reload screen."
        ),
        "expected": (
            "1. Awaiting Verification, Value Pending, Verified This Month and Dealers Submitting display "
            "for branch scope; Awaiting Verification equals pending row count; Value Pending equals sum of Amount column.\n"
            f"2. When systems are down (AC-22), validation message displays: \"{SYSTEM_DOWN_MSG}\""
        ),
        "remarks": "KPI Summary & System Down",
    },
    "TC_OPS2_003": {
        "ac": "AC-5, AC-23",
        "expected": (
            "1. Pending Verification is selected by default; Pending Verification, Verified, Rejected and All tabs available; "
            "explanatory caption is displayed.\n"
            "2. Where section data does not exist (AC-23), section shows empty state with grid column headers retained and no rows displayed."
        ),
        "remarks": "Queue Tabs & Empty Sections",
    },
    "TC_OPS2_008": {
        "ac": "AC-11, AC-13, AC-14, AC-20, AC-21",
        "steps": (
            "1. Approve a pending submission with optional note.\n"
            "2. Approve a submission whose UTR already exists on another verified submission (AC-20).\n"
            "3. Simulate LMS posting failure after verification (AC-21).\n"
            "4. Verify status outcomes."
        ),
        "expected": (
            "1. Approved submission moves to Verified; decision, user, timestamp and note recorded.\n"
            "2. Duplicate UTR does not block approval (AC-20).\n"
            "3. If LMS posting fails (AC-21), verification remains recorded, asset is not released and status returns to Pending Verification."
        ),
        "remarks": "Approval, UTR & LMS Failure",
    },
    "TC_OPS2_010": {
        "ac": "AC-16, AC-19",
        "steps": (
            "1. Record an approval or rejection for a pending submission.\n"
            "2. Observe KPI and queue update.\n"
            "3. Attempt to decide a submission already actioned by another Operations user (AC-19)."
        ),
        "expected": (
            "1. Decided submission excluded from Awaiting Verification and Value Pending; KPIs recalculate.\n"
            "2. Stale decision is rejected, row refreshed and current status displayed (AC-19)."
        ),
        "remarks": "Decision, KPI & Stale Action",
    },
    "TC_OPS3_001": {
        "title": "Verify access control and page header for Dealer Target",
        "ac": "AC-1, AC-2",
        "steps": (
            "Part A – Access denial (AC-1):\n"
            "1. Sign in as a user without dealer target entitlement.\n"
            "2. Attempt to access Ops Console – Dealer Target.\n\n"
            "Part B – Authorized access (AC-2):\n"
            "3. Sign in as authorised Operations user with dealer target entitlement.\n"
            "4. Open Dealer Target screen.\n"
            "5. Verify page header elements."
        ),
        "test_data": "Unauthorized Ops user; authorised Operations user with dealer target entitlement",
        "expected": (
            "Part A: Access is denied and module is not displayed.\n"
            "Part B: Ops Console badge, restricted-access caption, breadcrumb, page title and description, "
            "branch scope chip, and signed-in user name and role are displayed."
        ),
        "remarks": "Access & Header",
    },
    "TC_OPS3_002": {
        "ac": "AC-3, AC-4, AC-23, AC-24, AC-25",
        "steps": (
            "1. Open Dealer Target before any file upload.\n"
            "2. Verify KPI cards and upload panel empty state.\n"
            "3. Verify monetary formatting on KPI values.\n"
            "4. Simulate backend unavailability and observe empty sections."
        ),
        "expected": (
            "1. Active Dealers populated from dealer master; Dealers in File, Total File Amount and Active Period unset (AC-3, AC-4).\n"
            "2. Upload panel shows empty state with sample format prompt.\n"
            "3. Monetary values display in Indian numbering format with ₹ prefix (AC-23).\n"
            f"4. When systems are down (AC-24), validation message displays: \"{SYSTEM_DOWN_MSG}\"\n"
            "5. Where section data does not exist (AC-25), empty state shown with panel frame and KPI labels retained."
        ),
        "remarks": "KPIs, Empty State, Format & System Down",
    },
    "TC_OPS3_004": {
        "ac": "AC-7, AC-8, AC-21",
        "steps": (
            "1. Upload a valid dealer target file.\n"
            "2. Observe parsing loading state.\n"
            "3. Verify review state and summary strip update."
        ),
        "expected": (
            "1. Valid file is accepted, parsed and records displayed for review.\n"
            "2. Dealers in File, Total File Amount and Active Period update from parsed content (AC-8).\n"
            "3. During validation/parsing (AC-21), loading state displayed and Upload File/Submit actions disabled."
        ),
        "remarks": "Valid Upload, Review & Loading State",
    },
    "TC_OPS3_005": {
        "ac": "AC-9, AC-10, AC-11",
        "steps": (
            "1. Upload file with invalid type, size or column structure.\n"
            "2. Upload structurally valid file containing zero records.\n"
            "3. Observe Submit Target availability."
        ),
        "expected": (
            "1. Invalid type/size/columns rejected per AC-9/AC-10.\n"
            "2. Zero-record valid file: Submit Target is greyed out and user cannot submit (AC-11)."
        ),
        "remarks": "File Validation & Zero Records",
    },
    "TC_OPS3_006": {
        "ac": "AC-12, AC-13, AC-14, AC-17",
        "steps": (
            "1. Upload file with invalid dealer code or target/period format.\n"
            "2. Upload file with duplicate dealer rows for same period.\n"
            "3. Attempt submission."
        ),
        "expected": (
            "1. Invalid dealer code, target or period format rejected; commit prevented (AC-12, AC-13, AC-17).\n"
            "2. Duplicate dealer rows for same period are summed as total target for that dealer (AC-14)."
        ),
        "remarks": "Record Validation & Duplicate Sum",
    },
    "TC_OPS3_008": {
        "ac": "AC-15, AC-22",
        "expected": (
            "1. Successful submission commits dealer targets and records upload in audit (AC-15).\n"
            "2. When target file store is unavailable (AC-22), submission does not complete, no partial state "
            "is written and failure is raised to Operations."
        ),
        "remarks": "Commit, Audit & Store Failure",
    },
}


def find_row_by_tc_id(ws, tc_id: str, start_row: int = 4) -> int | None:
    for row in range(start_row, ws.max_row + 1):
        val = ws.cell(row=row, column=COL["id"]).value
        if val and str(val).strip() == tc_id:
            return row
    return None


def apply_update(ws, tc_id: str, fields: dict) -> None:
    row = find_row_by_tc_id(ws, tc_id)
    if row is None:
        raise KeyError(f"{tc_id} not found on sheet {ws.title}")
    for key, value in fields.items():
        ws.cell(row=row, column=COL[key], value=value)


def set_status_all(ws, start_row: int = 4) -> int:
    count = 0
    for row in range(start_row, ws.max_row + 1):
        tc_id = ws.cell(row=row, column=COL["id"]).value
        if tc_id and str(tc_id).strip():
            ws.cell(row=row, column=COL["status"], value="Not Executed")
            count += 1
    return count


def rebuild_all_test_cases(wb) -> None:
    ws_all = wb["All Test Cases"]
    if ws_all.max_row > 1:
        ws_all.delete_rows(2, ws_all.max_row - 1)
    row_idx = 2
    for sheet_name in STORY_SHEETS:
        ws = wb[sheet_name]
        for row in range(4, ws.max_row + 1):
            tc_id = ws.cell(row=row, column=COL["id"]).value
            if not tc_id or not str(tc_id).strip():
                continue
            for c in range(1, 13):
                ws_all.cell(row=row_idx, column=c, value=ws.cell(row=row, column=c).value)
            row_idx += 1


def main() -> None:
    wb = load_workbook(XLSX)
    total_status = 0
    update_count = 0

    for sheet_name in STORY_SHEETS:
        ws = wb[sheet_name]
        for tc_id, fields in UPDATES.items():
            if find_row_by_tc_id(ws, tc_id) is not None:
                apply_update(ws, tc_id, fields)
                update_count += 1
        total_status += set_status_all(ws)

    rebuild_all_test_cases(wb)

    if "Summary" in wb.sheetnames:
        ws_sum = wb["Summary"]
        for row in range(1, ws_sum.max_row + 1):
            val_a = ws_sum.cell(row=row, column=1).value
            if val_a and str(val_a).strip().upper() == "GRAND TOTAL":
                ws_sum.cell(row=row, column=3, value=149)

    out = XLSX
    try:
        wb.save(XLSX)
    except PermissionError:
        FALLBACK_XLSX.parent.mkdir(parents=True, exist_ok=True)
        wb.save(FALLBACK_XLSX)
        out = FALLBACK_XLSX
        print(f"WARNING: Could not write to {XLSX} (file may be open in Excel).")
        print(f"Saved to fallback: {FALLBACK_XLSX}")
    else:
        print(f"Saved: {XLSX}")
    print(f"Substantive updates applied: {update_count}")
    print(f"Status set to Not Executed: {total_status} rows across story sheets")
    print(f"All Test Cases rows: {wb['All Test Cases'].max_row - 1}")
    print(f"Output file: {out}")


if __name__ == "__main__":
    main()
